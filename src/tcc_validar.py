"""Validação humana da classificação: amostra cega + concordância (kappa de Cohen).

gerar  sorteia uma fração dos posts (padrão 12%), estratificada por fase,
       e gera uma planilha com SÓ id e texto, sem o rótulo da IA (para não
       ancorar o julgamento), com listas suspensas e uma aba com os critérios
       do protocolo. Faça a rotulagem ANTES de olhar o resultado da IA.
medir  compara com a classificação da IA: concordância, kappa de Cohen com
       IC 95% (bootstrap), matriz de confusão, precisão/recall por classe
       (rótulo humano como referência) e kappa por fase.

Uso:
  python src/tcc_validar.py gerar --caso tania_bulhoes
  # preencha dados/<caso>/validacao/<caso>_validacao_cega.xlsx
  python src/tcc_validar.py medir --caso tania_bulhoes
"""

from __future__ import annotations

import argparse
import random
import sys
from pathlib import Path

import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))
import tcc_protocolo as proto  # noqa: E402
from tcc_classificar import _sem_acento  # noqa: E402


def kappa(a: list[str], b: list[str]) -> float:
    n = len(a)
    if n == 0:
        return float("nan")
    cats = sorted(set(a) | set(b))
    po = sum(x == y for x, y in zip(a, b)) / n
    pe = sum((a.count(c) / n) * (b.count(c) / n) for c in cats)
    return 1.0 if pe == 1 else (po - pe) / (1 - pe)


def ic_bootstrap(a: list[str], b: list[str], reps: int = 2000, seed: int = 2025) -> tuple[float, float]:
    rng = random.Random(seed)
    n, ks = len(a), []
    for _ in range(reps):
        idx = [rng.randrange(n) for _ in range(n)]
        k = kappa([a[i] for i in idx], [b[i] for i in idx])
        if k == k:
            ks.append(k)
    ks.sort()
    return ks[int(0.025 * len(ks))], ks[int(0.975 * len(ks)) - 1]


def _md(df: pd.DataFrame, index: bool = True) -> str:
    """Tabela markdown sem depender do pacote tabulate."""
    t = df.reset_index() if index else df
    cab = [str(c) for c in t.columns]
    linhas = ["| " + " | ".join(cab) + " |", "|" + "---|" * len(cab)]
    linhas += ["| " + " | ".join(str(v) for v in r) + " |" for r in t.itertuples(index=False)]
    return "\n".join(linhas)


def interpretar(k: float) -> str:  # Landis & Koch (1977)
    for lim, txt in [(0, "pobre"), (0.2, "leve"), (0.4, "razoável"), (0.6, "moderada"),
                     (0.8, "substancial"), (1.01, "quase perfeita")]:
        if k <= lim:
            return txt
    return "quase perfeita"


def _paths(caso: str) -> tuple[Path, Path]:
    d = proto.tcc_dir(caso) / "validacao"
    d.mkdir(exist_ok=True)
    return d, d / f"{caso}_validacao_cega.xlsx"


def gerar(caso: str, frac: float, minimo: int) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.datavalidation import DataValidation

    cfg = proto.carregar(caso)
    seed = int(cfg["coleta"].get("seed", 2025))
    ccfg = cfg["classificacao"]
    posts = pd.read_csv(proto.tcc_dir(caso) / f"{caso}_posts.csv", sep=";", encoding="utf-8-sig", dtype=str)
    d, xlsx = _paths(caso)
    if xlsx.exists():
        sys.exit(f"{xlsx} já existe. Apague ou renomeie antes de sortear outra amostra.")

    # piso de `minimo` posts: com bases pequenas, 10-15% dá um IC do kappa largo demais
    frac = min(1.0, max(frac, minimo / len(posts)))
    partes = []
    for i, (fase, g) in enumerate(posts.groupby("fase")):
        n = max(1, round(len(g) * frac))
        partes.append(g.sample(n=min(n, len(g)), random_state=seed + i))
    am = pd.concat(partes).sample(frac=1, random_state=seed).reset_index(drop=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "rotular"
    ws.append(["id", "texto", "sentimento_humano", "eixo_humano", "ironia_humano", "observacao"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in am.itertuples():
        ws.append([r.id, r.texto_integral, "", "", "", ""])
    n = len(am) + 1
    for col, opcoes in [("C", ccfg["sentimentos"]), ("D", ccfg["eixos"]), ("E", ["sim", "nao"])]:
        dv = DataValidation(type="list", formula1='"' + ",".join(opcoes) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{n}")
    for col, w in zip("ABCDEF", [10, 90, 16, 28, 14, 30]):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"

    crit = wb.create_sheet("criterios")
    texto = Path(ccfg["prompt"]).read_text(encoding="utf-8").split("SAÍDA")[0]
    for ln in texto.splitlines():
        crit.append([ln])
    crit.column_dimensions["A"].width = 140
    wb.save(xlsx)

    am[["id", "fase"]].to_csv(d / f"{caso}_validacao_ids.csv", sep=";", index=False, encoding="utf-8-sig")
    print(f"{len(am)} posts ({len(am) / len(posts):.1%} da base) -> {xlsx}")
    print(am["fase"].value_counts().to_string())


def _tabela_metricas(h: list[str], m: list[str]) -> pd.DataFrame:
    linhas = []
    for c in sorted(set(h) | set(m)):
        tp = sum(x == c and y == c for x, y in zip(h, m))
        prec = tp / max(1, sum(y == c for y in m))
        rec = tp / max(1, sum(x == c for x in h))
        f1 = 2 * prec * rec / (prec + rec) if prec + rec else 0.0
        linhas.append({"classe": c, "n_humano": sum(x == c for x in h), "n_ia": sum(y == c for y in m),
                       "precisao": round(prec, 3), "recall": round(rec, 3), "f1": round(f1, 3)})
    return pd.DataFrame(linhas)


def medir(caso: str) -> None:
    out = proto.tcc_dir(caso)
    d, xlsx = _paths(caso)
    hum = pd.read_excel(xlsx, sheet_name="rotular", dtype=str).fillna("")
    ia = pd.read_csv(out / f"{caso}_classificacao.csv", sep=";", encoding="utf-8-sig", dtype=str)
    fases = pd.read_csv(d / f"{caso}_validacao_ids.csv", sep=";", encoding="utf-8-sig", dtype=str)
    for c in ["sentimento_humano", "eixo_humano", "ironia_humano"]:
        hum[c] = hum[c].map(_sem_acento)
    df = hum.merge(ia, on="id", how="inner").merge(fases, on="id", how="left")
    df = df[df["sentimento_humano"] != ""]
    if df.empty:
        sys.exit("Nenhum post rotulado ainda (coluna sentimento_humano vazia) ou sem classificação da IA.")

    rel = [f"# Validação humana x IA — {caso}\n",
           f"Posts rotulados e classificados: **{len(df)}**\n"]
    for nome, ch, ci in [("Sentimento", "sentimento_humano", "sentimento"),
                         ("Eixo temático", "eixo_humano", "eixo_tematico"),
                         ("Ironia", "ironia_humano", "ironia")]:
        sub = df[df[ch] != ""]
        if sub.empty:
            continue
        h, m = sub[ch].tolist(), sub[ci].tolist()
        k = kappa(h, m)
        lo, hi = ic_bootstrap(h, m)
        conc = sum(x == y for x, y in zip(h, m)) / len(h)
        rel += [f"\n## {nome} (n = {len(sub)})\n",
                f"- Concordância simples: {conc:.1%}",
                f"- Kappa de Cohen: **{k:.3f}** (IC 95% {lo:.3f} a {hi:.3f}) — concordância {interpretar(k)}\n"]
        if nome == "Sentimento":
            rel += ["### Matriz de confusão (linhas = humano, colunas = IA)\n",
                    _md(pd.crosstab(sub[ch], sub[ci])), "",
                    "### Por classe (humano como referência)\n",
                    _md(_tabela_metricas(h, m), index=False), "",
                    "### Kappa por fase\n"]
            for fase, g in sub.groupby("fase"):
                rel.append(f"- {fase}: kappa {kappa(g[ch].tolist(), g[ci].tolist()):.3f} (n = {len(g)})")
            div = sub[sub[ch] != sub[ci]][["id", ch, ci, "confianca", "justificativa"]]
            div.to_csv(d / f"{caso}_divergencias.csv", sep=";", index=False, encoding="utf-8-sig")
            rel.append(f"\nDivergências salvas em {caso}_divergencias.csv ({len(div)} posts), "
                       f"das quais {int((div['confianca'] == 'baixa').sum())} com confiança baixa da IA.")

    texto = "\n".join(rel)
    (d / f"{caso}_validacao_resultado.md").write_text(texto, encoding="utf-8")
    print(texto)


def main() -> None:
    ap = argparse.ArgumentParser(description="Validação humana cega + kappa de Cohen.")
    ap.add_argument("acao", choices=["gerar", "medir"])
    ap.add_argument("--caso", required=True, choices=proto.casos_disponiveis())
    ap.add_argument("--frac", type=float, default=0.12, help="fração da base (protocolo: 0.10 a 0.15)")
    ap.add_argument("--minimo", type=int, default=50, help="piso de posts na amostra (padrão 50)")
    a = ap.parse_args()
    gerar(a.caso, a.frac, a.minimo) if a.acao == "gerar" else medir(a.caso)


if __name__ == "__main__":
    main()
